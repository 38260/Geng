# -*- coding: utf-8 -*-
"""Deck 1 (开题汇报) — part A: slides 1-6.

All numbers are real, read from backend/data/gengv1.db and the documented
falsification experiment. Placeholders sit only where a human must fill in.
"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from gstyle import *  # noqa: F401,F403
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

HERE = os.path.dirname(os.path.abspath(__file__))
ASSET = os.path.join(HERE, "..", "assets")
ROOT = os.path.dirname(HERE)
COVER = os.path.join(ROOT, "docs", "pitch", "assets", "covers")
F = json.load(open(os.path.join(HERE, "deck_facts.json"), encoding="utf-8"))
TOTAL = 16
MASCOT = os.path.join(ROOT, "docs", "design", "asset-mascot-sign.png")

prs = new_deck()
N = [0]


def add(*, dark=False):
    s = blank(prs)
    canvas(s, dark=dark)
    N[0] += 1
    return s


def fin(s, note="Bilibili 真实采集 · 阈值集中写死在 backend/app/config/algorithms.py"):
    footer(s, N[0], TOTAL, note=note)


def blk(s, x, y, w, h, lines, *, warn=None, pad=1.0):
    return fit_text(s, x, y, w, h, lines, warn=warn or ("s%d" % N[0]), pad=pad)


def kv(s, x, y, w, h, big, label, sub=None, *, color=CORAL, bsize=30,
       fill=WHITE, line=LINE, bigfont=BODY):
    card(s, x, y, w, h, fill=fill, line=line)
    lines = [dict(text=big, size=bsize, color=color, bold=True, font=bigfont),
             dict(text=label, size=12, color=INK, bold=True, space_before=4)]
    if sub:
        lines.append(dict(text=sub, size=10, color=GREY, space_before=3,
                          line_spacing=1.2))
    blk(s, x + 0.26, y + 0.20, w - 0.52, h - 0.36, lines)


def thumb(s, name, x, y, w, h, score, stage, rank=None, tag=None):
    card(s, x, y, w, h, radius=0.06)
    img_h = h * 0.60
    p = os.path.join(COVER, name + ".jpg")
    if os.path.exists(p):
        pic_cover(s, p, x + 0.085, y + 0.085, w - 0.17, img_h)
    else:
        rect(s, x + 0.085, y + 0.085, w - 0.17, img_h, fill=NAVY_SOFT)
    if rank:
        b = rect(s, x + 0.20, y + 0.20, 0.40, 0.40,
                 fill=GOLD if rank == 1 else WHITE, shape=MSO_SHAPE.OVAL)
        b.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(b.text_frame, True, text=str(rank), size=13, bold=True,
             align=PP_ALIGN.CENTER, color=INK)
    ty = y + img_h + 0.20
    lines = [dict(text=name, size=13.5, color=INK, bold=True),
             dict(text="热度 %.1f · %s" % (score, stage), size=10, color=CORAL,
                  bold=True, space_before=3)]
    if tag:
        lines.append(dict(text=tag, size=9.5, color=GREY, space_before=2))
    blk(s, x + 0.18, ty, w - 0.36, y + h - ty - 0.10, lines)


def band(s, y, h, text, *, fill=CORAL_SOFT, color=INK, size=13, w=None,
         x=None, bold=True, extra=None, exsize=11):
    x = MARGIN if x is None else x
    w = CONTENT_W if w is None else w
    card(s, x, y, w, h, fill=fill, line=None)
    lines = [dict(text=text, size=size, color=color, bold=bold)]
    if extra:
        lines.append(dict(text=extra, size=exsize, color=GREY, space_before=4,
                          line_spacing=1.2))
    blk(s, x + 0.32, y + 0.15, w - 0.64, h - 0.28, lines)


# =========================================================================== #
# 1 — cover
# =========================================================================== #
s = add()
rect(s, 0, 0, SW, 5.60, fill=WHITE)
rect(s, 0, 0, 0.22, 5.60, fill=CORAL)
simple(s, MARGIN + 0.10, 0.72, 9.0, 0.32,
       [dict(text="数据科学课程 · 开题汇报", size=14, color=CORAL, bold=True)],
       warn="cover-kicker")
tb = textbox(s, MARGIN + 0.10, 1.06, 9.4, 1.05)
para(tb.text_frame, True, text="今天，赶什么梗？", size=54, color=INK,
     bold=True, font=DISPLAY, line_spacing=1.0)
rect(s, MARGIN + 0.10, 2.22, 1.90, 0.075, fill=CORAL)
simple(s, MARGIN + 0.10, 2.44, 9.0, 0.46,
       [dict(text="B 站网络梗热度与生命周期分析系统", size=19, color=NAVY,
             bold=True)], warn="cover-sub")
blk(s, MARGIN + 0.10, 3.06, 8.60, 2.32, [
    dict(text="开场先问一句", size=12, color=GREY),
    dict(text="「闪身步」9 月 21 日才有人解说，如今热度 76.8，全站第一；"
              "近 7 天头部播放 1,670 万。", size=15, color=INK, bold=True,
         space_before=6, line_spacing=1.3),
    dict(text="从「有人把它讲明白」到「火到不能再拖」，中间只有 5 天。",
         size=13.5, color=NAVY, space_before=7),
    dict(text="这 5 天，就是这个课题要研究的东西。", size=13.5, color=CORAL,
         bold=True, space_before=7),
], warn="cover-hook")
if os.path.exists(MASCOT):
    pic(s, MASCOT, 10.62, 0.98, h=2.05)
simple(s, 10.30, 3.18, 2.35, 0.30,
       [dict(text="今天，赶什么梗？", size=12, color=INK, bold=True,
             font=DISPLAY, align=PP_ALIGN.CENTER)], warn="cover-mascot-cap")
rect(s, 0, 5.60, SW, 1.90, fill=CANVAS)
info = card(s, 9.60, 5.86, 3.10, 1.60, radius=0.10)
blk(s, 9.86, 6.04, 2.58, 1.28, [
    dict(text="汇报人　＿＿＿＿＿＿", size=13, color=INK, bold=True),
    dict(text="学　号　＿＿＿＿＿＿", size=13, color=INK, space_before=6),
    dict(text="指导教师　＿＿＿＿＿＿", size=13, color=INK, space_before=6),
    dict(text="＿＿＿＿ 年 ＿＿ 月 ＿＿ 日", size=10.5, color=GREY,
         space_before=7),
], warn="cover-info")
band(s, 5.86, 1.60, "先把话说在前面：本阶段交付的是「一条跑通的真实数据链路 + 一套可复算的测量口径」，"
                    "不是一堆已经成立的结论。",
     fill=CORAL_SOFT, size=13.5, w=8.60, x=MARGIN,
     extra="里面有一个我自己先立、后被数据推翻的实验（第 11 页）——"
           "我认为那才是这个课题最有价值的部分。")
fin(s, note="数据截至 2026-09-28，真实抓取；演示梗单独标注、不进真实榜单")
notes(s, "开场不要念标题。\n"
         "直接问：有谁听过「闪身步」？——它是 9 月 21 日才被解说、9 月 26 日就冲到榜首的梗。\n"
         "然后主动交底：现在给的是链路和口径，不是结论。把预期压低，后面被追问时不会翻车。\n"
         "数字来源：hotness_snapshots 里「闪身步」= 76.8（爆发期，还来得及，置信度 0.95），"
         "近 7 天头部播放 16,705,446。")

# =========================================================================== #
# 2 — hook
# =========================================================================== #
s = add()
header(s, "先测一下", "这五个梗，认识三个以上的举手",
       sub="不是我编的题——它们就是此刻 B 站真实热榜前五名。封面是真实视频封面，数字是库里现算的。")
names = [("闪身步", 76.8, "爆发期", "爆发期 · 置信度 0.95"),
         ("耍起", 67.2, "上升期", "上升期 · 还来得及"),
         ("闹吃vs古振兴", 54.6, "上升期", "上升期 · 慎赶"),
         ("野生狗奶", 47.5, "上升期", "双 UP 认证 · 慎赶"),
         ("大狗叫", 47.0, "数据不足", "近 7 天只观测到几天")]
x = MARGIN
NCOL = 5
GAP = 0.19
CW = (CONTENT_W - (NCOL - 1) * GAP) / NCOL
for i, (n, sc, st, tg) in enumerate(names):
    thumb(s, n, x, BODY_TOP, CW, 2.98, sc, st, rank=i + 1, tag=tg)
    x += CW + GAP
band(s, BODY_TOP + 3.16, 1.34,
     "认不全太正常了：梗的活跃窗口一般就两三周。难受的地方不在这——"
     "你听说了、点开创作中心、脚本写完发出去，它已经不火了。",
     fill=CORAL_SOFT, size=14,
     extra="注意第 5 名「大狗叫」：热度 47.0 排进前十，但近 7 天真正观测到的天数太少，"
           "算法拒绝给阶段与赶梗结论，标成「数据不足」。宁可说不知道，也不编一个阶段。")
fin(s)
notes(s, "现场互动 10 秒，别恋战。\n"
         "重点落在最后那句：第 5 名被标成「数据不足」。这是全课题的立场——"
         "拿接口空洞判「正在退潮」，是这类产品最不该犯的错。\n"
         "五个名字与分数都来自 lifecycle_snapshots / hotness_snapshots，可现场查库。")

# =========================================================================== #
# 3 — why it hurts
# =========================================================================== #
s = add()
header(s, "问题在哪", "三家人在问同一句话：现在做，还来得及吗",
       sub="他们都缺同一个东西——不是排名，是刻度。")
rows = [
    ("一个 3 万粉的 UP 主", "「这个梗我上周就刷到了，现在拍是不是凉了？」",
     "他有手感，但没有刻度：同一个梗，他每次判断都不一样。", BLUE, BLUE_SOFT),
    ("品牌社媒运营", "「热搜上这个词，我们今晚要不要跟？」",
     "跟的是词，不是梗：词还挂在榜上，梗可能三天前就过气了。", GOLD_DEEP, RGBColor(0xFF, 0xF7, 0xE3)),
    ("我室友（也是目标用户）", "「为什么我每次知道一个梗，它都已经不好玩了？」",
     "因为他拿到的永远是排名，不是时间。", CORAL, CORAL_SOFT),
]
y = BODY_TOP
RH = 1.10
for title, quote, note, ac, bg in rows:
    card(s, MARGIN, y, CONTENT_W, RH, fill=WHITE, line=LINE)
    rect(s, MARGIN, y, 0.055, RH, fill=ac)
    blk(s, MARGIN + 0.34, y + 0.15, 3.10, 0.84, [
        dict(text=title, size=13, color=ac, bold=True),
        dict(text=note, size=10.5, color=GREY, space_before=5,
             line_spacing=1.2)])
    blk(s, MARGIN + 3.70, y + 0.22, CONTENT_W - 4.10, 0.68, [
        dict(text=quote, size=16, color=INK, bold=True, font=DISPLAY)])
    y += RH + 0.15
band(s, y + 0.03, 1.40,
     "现有工具都能回答「此刻什么排在前面」，回答不了三个更值钱的问题——"
     "有多火（一个跨梗可比的数）、走到哪一段（在冲、在顶、还是在崩）、现在赶还来不来得及。",
     fill=NAVY_SOFT, color=NAVY, size=13.5,
     extra="微博热搜与平台热榜是存量快照：一个梗崩盘当天照样挂在榜上；"
           "百度指数度量的是搜索词，不是梗——「松弛感」这类口语词会混进大量无关检索。")
fin(s)
notes(s, "这一页讲「需求真实存在」，用三句原话，不要讲成抽象的市场分析。\n"
         "关键一句：他们要的不是排名，是刻度。\n"
         "如果老师问「这跟热搜有什么不同」：热搜是存量快照 + 关键词，"
         "我们是逐日时间序列 + 梗级聚合，还能给出剩余窗口判断。")

# =========================================================================== #
# 4 — RQ
# =========================================================================== #
s = add()
header(s, "研究问题", "把「还赶不赶得上」拆成三个能被回答的问题",
       sub="每个问题都对应一个已实现的算法模块；不能量化的一律不做。")
rqs = [
    ("RQ1", "一个梗现在到底有多火？", CORAL, CORAL_SOFT,
     ["输出 0-100 热度指数：五因子加权，绝对量先做对数区间归一化",
      "同一个数可以横向比两个梗，也可以纵向比一个梗的两周",
      "实现：backend/app/analytics/hotness.py"]),
    ("RQ2", "它走到了生命周期的哪一段？", NAVY, NAVY_SOFT,
     ["六个阶段：萌芽 / 上升 / 爆发 / 平稳 / 退潮 / 过气",
      "只用四个可观测量判定：增长率、距峰值差、连续下滑天数、活跃度",
      "实现：backend/app/analytics/lifecycle.py（规则自上而下匹配）"]),
    ("RQ3", "现在赶还来得及吗？", GOLD_DEEP, RGBColor(0xFF, 0xF7, 0xE3),
     ["三态输出：还来得及 / 慎赶 / 你来晚了，附带置信度与一句理由",
      "由 RQ1 + RQ2 结果经判定规则导出，不是独立模型猜的",
      "实现：backend/app/analytics/catch_up.py"]),
]
x = MARGIN
NCOL = 3
CW = (CONTENT_W - (NCOL - 1) * 0.17) / NCOL
for code, title, ac, bg, bullets in rqs:
    card(s, x, BODY_TOP, CW, 3.36, fill=WHITE, line=LINE)
    rect(s, x, BODY_TOP, CW, 0.075, fill=ac)
    tag = card(s, x + 0.28, BODY_TOP + 0.26, 0.80, 0.36, fill=bg, line=None,
               radius=0.26)
    tag.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tag.text_frame, True, text=code, size=12, color=ac, bold=True,
         align=PP_ALIGN.CENTER)
    lines = [dict(text=title, size=15.5, color=INK, bold=True, space_before=6,
                  line_spacing=1.18)]
    for b in bullets:
        lines.append(dict(text="· " + b, size=10.5, color=GREY, space_before=6,
                          line_spacing=1.24))
    blk(s, x + 0.28, BODY_TOP + 0.78, CW - 0.56, 2.36, lines)
    x += CW + 0.17
band(s, BODY_TOP + 3.56, 1.42,
     "明确不做的三件事，写在幻灯片上，也写在代码里：",
     fill=NAVY_SOFT, color=NAVY, size=13,
     extra="① 不做未来数值预测，只描述已经发生的数据；"
           "② 不让大模型修改任何事实或判定——它只能把算法结论翻译成人话，"
           "越界表述（「预计」「一定会爆」）在服务层直接丢弃；"
           "③ 不做多平台、不做用户体系，先把 B 站这一把尺子做准。")
fin(s)
notes(s, "这一页是全场的技术骨架，讲慢一点。\n"
         "三个问题的顺序有讲究：先有可比的数，才有阶段，才有「赶不赶」。"
         "反过来做（先问赶不赶）就是拍脑袋。\n"
         "老师最容易问「为什么不用深度学习预测」——答案在 RQ3 那句："
         "本课题的输出是判定问题，不是外推问题；而且样本量撑不起深度模型"
         "（84 条梗记录、2,209 行日统计）。")

# =========================================================================== #
# 5 — literature gap
# =========================================================================== #
s = add()
header(s, "文献位置", "三条线各自都做到了，缺的是把它们接起来的那个人",
       sub="不是没人研究，是没人把「梗」当成一个能跨创作者聚合、还能复算的时间序列对象。")
cols = [
    ("模因与网络梗研究", NAVY, "定义、分类、生成机制、社会心态", 
     "Dawkins 1976 · Shifman 2012/2013 · Bauckhage ICWSM 2011 · 谢朝群 2007 · 薛一飞 2024",
     "几乎不给出可复算的时间度量"),
    ("注意力与话题生命周期", BLUE, "曲线形状可分类、半衰期在缩短、阶段可划分",
     "Downs 1972 · Crane & Sornette PNAS 2008 · Kleinberg KDD 2002 · Yang & Leskovec WSDM 2011 · Lorenz-Spreen 2019",
     "对象是新闻 / 议题 / 搜索词，不是「梗」"),
    ("流行度量化与预测", GOLD_DEEP, "早期增速可预测、热度寿命可建模",
     "Szabo & Huberman CACM 2010 · Tan BMSB 2014 · Li 2016 · Xu CIKM 2025",
     "落到单条内容，未做跨创作者的「梗」级聚合"),
]
x = MARGIN
NCOL = 3
CW = (CONTENT_W - (NCOL - 1) * 0.17) / NCOL
for title, ac, got, cites, gap in cols:
    card(s, x, BODY_TOP, CW, 3.20, fill=WHITE, line=LINE)
    rect(s, x, BODY_TOP, CW, 0.075, fill=ac)
    lines = [
        dict(text=title, size=14.5, color=INK, bold=True),
        dict(text="做到了什么", size=10, color=ac, bold=True, space_before=12),
        dict(text=got, size=11.5, color=INK, space_before=3, line_spacing=1.24),
        dict(text="代表工作", size=10, color=ac, bold=True, space_before=12),
        dict(text=cites, size=9.5, color=GREY, line_spacing=1.22),
        dict(text="留下的缺口", size=10, color=CORAL, bold=True, space_before=12),
        dict(text=gap, size=11.5, color=INK, space_before=3, line_spacing=1.24),
    ]
    blk(s, x + 0.28, BODY_TOP + 0.28, CW - 0.56, 2.68, lines)
    x += CW + 0.17
band(s, BODY_TOP + 3.40, 1.58,
     "本课题的位置：把三条线接成一条能跑起来的链路。",
     fill=CORAL_SOFT, size=14,
     extra="用两位解说 UP 主的真实投稿证明「这是个真梗」→ 用逐日头部内容抽样把它变成时间序列 → "
           "用五因子指数与规则引擎给出热度、阶段与赶梗判断 → 把全部口径与数据新鲜度"
           "直接印在产品界面上。文献核验记录（DOI / CNKI）在 docs/research/related-literature.md。")
fin(s)
notes(s, "这一页不要逐条念文献。只讲结构：三条线、三个缺口、一个位置。\n"
         "如果老师是传播学背景，重点接 Lorenz-Spreen 2019（注意力半衰期在缩短）——"
         "它说明「还剩多久」比「有多热」更值钱，直接支撑 RQ3。\n"
         "如果老师是数据挖掘背景，重点接 Szabo & Huberman 2010（早期增速可预测）——"
         "它支撑我们把增长权重给到 0.25。\n"
         "所有 DOI 都核验过，被质疑引用可当场翻 docs/research/related-literature.md。")

# =========================================================================== #
# 6 — feasibility: the pipeline already runs
# =========================================================================== #
s = add()
header(s, "可行性：不是设想", "这条链路已经跑通了，今天的数字都是它算出来的",
       sub="先把「能不能做」回答掉，后面才轮得到「做得对不对」。",
       tag="已可运行")
# pipeline
steps = [("B 站接口", NAVY), ("逐日采集", NAVY), ("清洗 + 梗匹配", NAVY),
         ("发现层准入", CORAL), ("日粒度聚合", NAVY), ("热度指数", CORAL),
         ("生命周期", CORAL), ("赶梗判断", CORAL), ("FastAPI", NAVY),
         ("React", NAVY)]
x = MARGIN
NCOL = 10
CW = (CONTENT_W - (NCOL - 1) * 0.045) / NCOL
for i, (t, c) in enumerate(steps):
    card(s, x, BODY_TOP, CW, 0.86, fill=WHITE if c == NAVY else CORAL_SOFT,
         line=LINE if c == NAVY else None, radius=0.14)
    blk(s, x + 0.06, BODY_TOP + 0.10, CW - 0.12, 0.66,
        [dict(text=t, size=11, color=c if c == NAVY else CORAL, bold=True,
              align=PP_ALIGN.CENTER, line_spacing=1.1)])
    x += CW + 0.045
simple(s, MARGIN, BODY_TOP + 0.98, CONTENT_W, 0.26,
       [dict(text="数据负责证明　·　算法负责判断　·　大模型负责解释　·　界面负责让人想看",
             size=11, color=GREY, align=PP_ALIGN.CENTER)])
kvs = [
    ("69", "后端 Python 模块", "10,409 行；另有 22 个测试文件 4,507 行", CORAL),
    ("279", "后端测试全绿", "假客户端 + 内存库，全部离线可复跑", NAVY),
    ("125", "接口冒烟用例", "对着运行中的真实后端逐个打过去", NAVY),
    ("31", "前端 TS/TSX 文件", "4,666 行，Vite 构建产物 204KB（gzip 66KB）", NAVY),
    ("11", "秒", "依赖就绪时 start.bat 一键起前后端；空库自动灌演示数据", GOLD_DEEP),
]
x = MARGIN
NCOL = 5
GAP = 0.19
CW = (CONTENT_W - (NCOL - 1) * GAP) / NCOL
for big, label, sub, c in kvs:
    kv(s, x, BODY_TOP + 1.34, CW, 1.28, big, label, sub, color=c, bsize=26)
    x += CW + GAP
band(s, BODY_TOP + 2.72, 0.86,
     "数据可得性也验证过了，不是假设：WBI 签名后匿名可用的只有搜索与视频详情接口。",
     fill=NAVY_SOFT, color=NAVY, size=12.5,
     extra="UP 主空间接口会被风控（-352 / 412），字幕轨与站内 AI 总结端点匿名一律拿不到"
           "（实测 code=0 但 subtitles 为空、总结接口直接 -101 未登录）。"
           "这两件事都如实写进了产品的「数据透明度」，没有拿旧证据凑数。")
fin(s)
notes(s, "这一页的作用是让老师放心：这不是一个 PPT 项目。\n"
         "建议现场直接演示 start.bat（约 11 秒）。\n"
         "数字出处：LOC 与测试数用仓库实测统计；279 项由 python -m pytest 收集得出；"
         "125 项来自 scripts/smoke_api.py。\n"
         "最后那条限制很重要，主动说：字幕这一档要 Cookie 才有，"
         "所以界面上「字幕原文摘录」现在可能是空的——那是如实标注，不是坏掉。")

prs.save(os.path.join(HERE, "deck_proposal_A.pptx"))
print("part A ok, slides:", N[0])
