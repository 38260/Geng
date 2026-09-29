# -*- coding: utf-8 -*-
"""Audit a deck: recompute every text box's required height with the same
metric model used to build it, and report any box whose text overflows."""
import sys, os, json
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from gstyle import measure_lines, BODY, DISPLAY, FOOT_Y, SH, SW

from pptx import Presentation
from pptx.util import Emu

EMU = 914400.0

def audit(path):
    prs = Presentation(path)
    print("=" * 78)
    print(path, "| slides:", len(prs.slides),
          "| %.3f x %.3f in" % (prs.slide_width / EMU, prs.slide_height / EMU))
    bad = 0
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            txt = sh.text_frame.text
            if not txt.strip():
                continue
            w = sh.width / EMU
            h = sh.height / EMU
            lines = []
            for p in sh.text_frame.paragraphs:
                runs = [r for r in p.runs if r.text]
                if not runs:
                    continue
                text = "".join(r.text for r in runs)
                r0 = runs[0]
                size = r0.font.size.pt if r0.font.size else 18
                font = r0.font.name or BODY
                ls = p.line_spacing or 1.22
                lines.append(dict(text=text, size=size, font=font,
                                  line_spacing=ls,
                                  space_before=(p.space_before.pt if p.space_before else 0),
                                  space_after=(p.space_after.pt if p.space_after else 0)))
            need = measure_lines(lines, w)
            bottom = (sh.top / EMU) + h
            problems = []
            if need > h + 0.02:
                problems.append("text %.2fin > box %.2fin" % (need, h))
            if bottom > SH - 0.02:
                problems.append("box bottom %.2fin beyond slide" % bottom)
            if (sh.left / EMU) + w > SW - 0.02:
                problems.append("box right %.2fin beyond slide" % ((sh.left / EMU) + w))
            if problems:
                bad += 1
                print("  slide %-2d %-22s %s" % (i, sh.shape_type, "; ".join(problems)))
                print("        text: %s" % txt.replace("\n", " / ")[:80])
    print("  --- %d problem shape(s)" % bad)
    return bad

total = 0
for p in sys.argv[1:]:
    total += audit(p)
print()
print("TOTAL PROBLEMS:", total)
